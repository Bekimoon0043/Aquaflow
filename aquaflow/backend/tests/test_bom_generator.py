"""
Contract tests for the automatic BOM generator (AQ-S2-009).

Asserts quantities are aggregated automatically, supplier references are attached
from the catalog, and the Excel export is a valid .xlsx workbook.
"""
import io
import zipfile
from xml.dom import minidom

from services.bom_generator import (
    BOMGenerator, generate_bom, generate_from_network, build_xlsx, _col_ref,
)


def _model():
    return {
        "project": "P", "currency": "USD",
        "equipment": [{"type": "pump", "part_number": "pump", "quantity": 1},
                      {"type": "pump", "part_number": "pump", "quantity": 1}],
        "nozzles": [{"model": "AF-GEYSER-25"}, {"model": "AF-GEYSER-25"}, {"model": "AF-GEYSER-25"}],
        "pipes": [{"diameter": 0.1, "length": 12.0, "material": "PVC"},
                  {"diameter": 0.1, "length": 8.0, "material": "PVC"}],
        "valves": [{"type": "control_valve", "size": "DN100", "quantity": 2}],
        "instruments": [{"type": "PT", "quantity": 1}],
    }


def test_quantities_are_aggregated():
    gen = BOMGenerator().generate(_model())
    pump = next(l for l in gen.lines if l.category == "pump")
    assert pump.quantity == 2
    geyser = next(l for l in gen.lines if "GEYSER" in l.part_number)
    assert geyser.quantity == 3
    pipe = next(l for l in gen.lines if l.category == "pipe")
    assert abs(pipe.quantity - 20.0) < 1e-9   # 12 + 8 summed


def test_supplier_references_attached():
    gen = BOMGenerator().generate(_model())
    pump = next(l for l in gen.lines if l.category == "pump")
    assert pump.manufacturer == "Grundfos"
    assert pump.supplier == "Grundfos Distribution"
    assert pump.unit_cost == 1850.0
    assert pump.total_cost == 3700.0


def test_total_cost_rollup():
    gen = BOMGenerator().generate(_model())
    hand = 2 * 1850.0 + 3 * 95.0 + 20.0 * 15.8 + 2 * 420.0 + 1 * 240.0
    assert abs(gen.total_cost() - round(hand, 2)) < 1e-6


def test_col_ref():
    assert _col_ref(0) == "A"
    assert _col_ref(25) == "Z"
    assert _col_ref(26) == "AA"


def test_xlsx_is_valid_package():
    xlsx = BOMGenerator().generate(_model()).to_xlsx_bytes()
    zf = zipfile.ZipFile(io.BytesIO(xlsx))
    required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
                "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml"}
    assert required <= set(zf.namelist())
    minidom.parseString(zf.read("xl/worksheets/sheet1.xml"))   # well-formed


def test_xlsx_reads_back_with_openpyxl():
    import openpyxl
    xlsx = BOMGenerator().generate(_model()).to_xlsx_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header[0] == "Item" and "Supplier" in header
    pump_rows = [r for r in ws.iter_rows(values_only=True) if r[1] == "pump"]
    assert pump_rows and pump_rows[0][3] == 2         # quantity reads back as number
    total_rows = [r for r in ws.iter_rows(values_only=True) if r[2] == "TOTAL"]
    assert total_rows and isinstance(total_rows[0][9], (int, float))


def test_build_xlsx_escapes_and_types():
    xlsx = build_xlsx("S", ["A&B", "n"], [["<x>", 5], ["y", 6.5]])
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).active
    assert ws["A1"].value == "A&B"        # XML-escaped round-trip
    assert ws["A2"].value == "<x>"
    assert ws["B2"].value == 5            # numeric preserved
    assert ws["B3"].value == 6.5


def test_generate_from_network_sums_pipes():
    out = generate_from_network({
        "nodes": [{"id": "PUMP_OUT", "fixed_head": 55},
                  {"id": "J1", "demand": 0.01, "nozzle_model": "AF-GEYSER-25"}],
        "pipes": [{"id": "s", "from": "PUMP_OUT", "to": "J1", "diameter": 0.1, "length": 10, "material": "PVC"},
                  {"id": "m", "from": "J1", "to": "J1", "diameter": 0.1, "length": 5, "material": "PVC"}],
    })
    pipe = next(l for l in out["lines"] if l["category"] == "pipe")
    assert abs(pipe["quantity"] - 15.0) < 1e-9
    assert out["summary"]["line_count"] >= 2
